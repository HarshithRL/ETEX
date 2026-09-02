from __future__ import annotations

"""Apply patch ops to an ArtifactSpec, and convert parsed documents back to spec.

Office formats are patched in place by the writers. PDFs use this module:
parse → document_to_spec → apply_ops_to_spec → write.
"""

import re
from typing import Any

from openpyxl.utils import column_index_from_string, coordinate_from_string

from .exceptions import ArtifactPatchError
from .models import ArtifactDocument, ArtifactType, BlockType, ContentBlock, SourceLocation
from .spec import (
    ArtifactSpec,
    PatchOp,
    SpecBlock,
    SpecSheet,
    SpecSlide,
    coerce_block,
    table_as_text,
)

_LIST_PREFIX = re.compile(r"^\s*([•●○▪►▪]|[-*+]|\d+[.)]|[A-Za-z][.)])\s+")


def document_to_spec(document: ArtifactDocument) -> ArtifactSpec:
    spec = ArtifactSpec(
        artifact_type=document.artifact_type,
        filename=document.metadata.filename,
        title=document.metadata.title,
        author=document.metadata.author,
    )
    if document.artifact_type == ArtifactType.PPT:
        spec.slides = _slides_from_blocks(document.blocks)
    elif document.artifact_type == ArtifactType.EXCEL:
        spec.sheets = [
            SpecSheet(
                name=block.location.sheet or f"Sheet{index + 1}",
                rows=[list(row) for row in (block.table or [])],
                block_id=block.block_id,
            )
            for index, block in enumerate(document.blocks)
        ]
        if not spec.sheets:
            spec.sheets = [SpecSheet(name="Sheet1")]
    else:
        spec.blocks = [content_to_spec_block(block) for block in document.blocks]
    return spec


def content_to_spec_block(block: ContentBlock) -> SpecBlock:
    items = None
    if block.type == BlockType.LIST:
        items = [_strip_list_prefix(line) for line in block.text.splitlines() if line.strip()]
    return SpecBlock(
        type=block.type,
        text=block.text,
        table=[list(row) for row in block.table] if block.table is not None else None,
        level=block.level,
        block_id=block.block_id,
        items=items,
        extra=dict(block.extra),
        location=block.location,
    )


def apply_ops_to_spec(spec: ArtifactSpec, ops: list[PatchOp]) -> None:
    for op in ops:
        apply_op_to_spec(spec, op)


def apply_op_to_spec(spec: ArtifactSpec, op: PatchOp) -> None:
    if op.op == "set_title":
        spec.title = op.title or op.text
        if spec.title is None:
            raise ArtifactPatchError("set_title requires 'title' or 'text'.")
        return
    if op.op in {"replace_text", "replace_shape_text"}:
        target = find_spec_block(spec, op)
        target.text = require_text(op)
        if target.type == BlockType.LIST:
            target.items = [line.strip() for line in target.text.splitlines() if line.strip()]
        return
    if op.op == "replace_table":
        table = require_table(op)
        kind, container, index, target = find_spec_target(spec, op)
        if kind == "sheet":
            target.rows = [list(row) for row in table]
            return
        target.table = [list(row) for row in table]
        target.type = BlockType.TABLE
        target.text = table_as_text(table)
        return
    if op.op == "set_cell":
        _apply_set_cell_spec(spec, op)
        return
    if op.op == "insert_block":
        _insert_spec_block(spec, op)
        return
    if op.op == "delete_block":
        kind, container, index, _target = find_spec_target(spec, op)
        container.pop(index)
        return
    if op.op == "insert_slide":
        if spec.artifact_type != ArtifactType.PPT:
            raise ArtifactPatchError("insert_slide is only valid for ppt artifacts.")
        if op.slide is None:
            raise ArtifactPatchError("insert_slide requires a 'slide' payload.")
        if spec.slides is None:
            spec.slides = []
        spec.slides.append(op.slide)
        return
    raise ArtifactPatchError(f"Unknown patch op {op.op!r}.")


def find_content_block(blocks: list[ContentBlock], op: PatchOp) -> ContentBlock:
    require_block_target(op)
    if op.block_id:
        for block in blocks:
            if block.block_id == op.block_id or block.id == op.block_id:
                return block
        raise ArtifactPatchError(f"No block with id {op.block_id!r}.")
    matches = [block for block in blocks if location_matches(block.location, op.location)]
    if not matches:
        raise ArtifactPatchError("No block matches the given location.")
    if len(matches) > 1:
        raise ArtifactPatchError("Location matches multiple blocks; pass block_id.")
    return matches[0]


def find_spec_block(spec: ArtifactSpec, op: PatchOp) -> SpecBlock:
    kind, _container, _index, target = find_spec_target(spec, op)
    if kind != "block":
        raise ArtifactPatchError(f"{op.op} expected a content block, not a sheet.")
    return target


def find_spec_target(spec: ArtifactSpec, op: PatchOp) -> tuple[str, list[Any], int, Any]:
    require_block_target(op)
    if op.block_id:
        found = _spec_by_id(spec, op.block_id)
        if found is None:
            raise ArtifactPatchError(f"No block with id {op.block_id!r}.")
        return found
    if op.location is not None:
        matches = []
        for kind, container, index, target in _iter_spec_targets(spec):
            loc = getattr(target, "location", None)
            if loc is not None and location_matches(loc, op.location):
                matches.append((kind, container, index, target))
            elif kind == "sheet" and op.location.sheet and target.name == op.location.sheet:
                matches.append((kind, container, index, target))
        if not matches:
            raise ArtifactPatchError("No block matches the given location.")
        if len(matches) > 1:
            raise ArtifactPatchError("Location matches multiple blocks; pass block_id.")
        return matches[0]
    raise ArtifactPatchError(f"{op.op} requires block_id or location.")


def location_matches(loc: SourceLocation, query: SourceLocation | None) -> bool:
    if query is None:
        return False
    specified = False
    for name in (
        "page",
        "slide",
        "sheet",
        "paragraph_index",
        "table_index",
        "row",
        "col",
        "cell_range",
        "shape_id",
    ):
        wanted = getattr(query, name)
        if wanted is None or wanted == "":
            continue
        specified = True
        if getattr(loc, name) != wanted:
            return False
    return specified


def require_block_target(op: PatchOp) -> None:
    if op.op in {"set_title", "insert_slide", "set_cell", "insert_block"}:
        return
    if not op.block_id and op.location is None:
        raise ArtifactPatchError(f"{op.op} requires block_id or location.")


def require_text(op: PatchOp) -> str:
    if op.text is None:
        raise ArtifactPatchError(f"{op.op} requires 'text'.")
    return op.text


def require_table(op: PatchOp) -> list[list[str]]:
    if op.table is None:
        raise ArtifactPatchError(f"{op.op} requires 'table'.")
    return op.table


def cell_value(op: PatchOp) -> Any:
    if op.value is not None:
        return op.value
    if op.text is not None:
        return op.text
    raise ArtifactPatchError("set_cell requires 'value' (or 'text').")


def parse_cell_ref(op: PatchOp) -> tuple[int, int]:
    """Return 1-based (row, col) for a set_cell op."""
    if op.cell_range:
        coord = op.cell_range.split(":", 1)[0].strip().upper()
        try:
            col_letter, row = coordinate_from_string(coord)
            return int(row), int(column_index_from_string(col_letter))
        except Exception as exc:
            raise ArtifactPatchError(f"Invalid cell_range {op.cell_range!r}.") from exc
    if op.row is not None and op.col is not None:
        if op.row < 1 or op.col < 1:
            raise ArtifactPatchError("row and col are 1-based Excel coordinates.")
        return op.row, op.col
    raise ArtifactPatchError("set_cell requires cell_range or row+col.")


def set_cell_in_rows(rows: list[list[str]], row: int, col: int, value: Any) -> list[list[str]]:
    while len(rows) < row:
        rows.append([])
    line = rows[row - 1]
    while len(line) < col:
        line.append("")
    line[col - 1] = "" if value is None else str(value)
    return rows


class LocationTracker:
    """Mutable view of parsed blocks so insert/delete can shift native indices."""

    def __init__(self, blocks: list[ContentBlock]) -> None:
        self.blocks = list(blocks)

    def find(self, op: PatchOp) -> ContentBlock:
        return find_content_block(self.blocks, op)

    def find_anchor(self, op: PatchOp) -> ContentBlock:
        anchor_id = op.after_block_id or op.block_id
        if not anchor_id and op.location is None:
            raise ArtifactPatchError("insert_block requires after_block_id, block_id, or location.")
        probe = PatchOp(op="insert_block", block_id=anchor_id, location=op.location)
        return find_content_block(self.blocks, probe)

    def shift_paragraphs(self, after_index: int, delta: int) -> None:
        for block in self.blocks:
            idx = block.location.paragraph_index
            if idx is not None and idx > after_index:
                block.location.paragraph_index = idx + delta

    def shift_tables(self, after_index: int, delta: int) -> None:
        for block in self.blocks:
            idx = block.location.table_index
            if idx is not None and idx > after_index:
                block.location.table_index = idx + delta

    def remove(self, block: ContentBlock) -> None:
        self.blocks = [item for item in self.blocks if item is not block]


def _apply_set_cell_spec(spec: ArtifactSpec, op: PatchOp) -> None:
    sheet = _resolve_spec_sheet(spec, op)
    row, col = parse_cell_ref(op)
    sheet.rows = set_cell_in_rows(list(sheet.rows), row, col, cell_value(op))


def _resolve_spec_sheet(spec: ArtifactSpec, op: PatchOp) -> SpecSheet:
    if spec.sheets is None:
        raise ArtifactPatchError("set_cell requires an excel spec with sheets.")
    if op.block_id:
        found = _spec_by_id(spec, op.block_id)
        if found is None:
            raise ArtifactPatchError(f"No block with id {op.block_id!r}.")
        kind, _container, _index, target = found
        if kind == "sheet":
            return target
        raise ArtifactPatchError("set_cell block_id must refer to a sheet table.")
    name = op.sheet or (op.location.sheet if op.location else None)
    if name:
        for sheet in spec.sheets:
            if sheet.name == name:
                return sheet
        raise ArtifactPatchError(f"Sheet not found: {name}.")
    if len(spec.sheets) == 1:
        return spec.sheets[0]
    raise ArtifactPatchError("set_cell requires sheet, block_id, or a single-sheet workbook.")


def _insert_spec_block(spec: ArtifactSpec, op: PatchOp) -> None:
    if op.block is None:
        raise ArtifactPatchError("insert_block requires a 'block' payload.")
    new_block = coerce_block(op.block)
    anchor_id = op.after_block_id or op.block_id
    if not anchor_id and op.location is None:
        spec.blocks.append(new_block)
        return
    probe = PatchOp(op="insert_block", block_id=anchor_id, location=op.location)
    kind, container, index, _target = find_spec_target(spec, probe)
    if kind != "block":
        raise ArtifactPatchError("insert_block after a sheet is not supported; insert a new sheet via insert_block on excel in-place.")
    container.insert(index + 1, new_block)


def _spec_by_id(spec: ArtifactSpec, block_id: str) -> tuple[str, list[Any], int, Any] | None:
    for kind, container, index, target in _iter_spec_targets(spec):
        target_id = getattr(target, "block_id", "") or ""
        if target_id == block_id:
            return kind, container, index, target
    return None


def _iter_spec_targets(spec: ArtifactSpec):
    for index, block in enumerate(spec.blocks):
        yield "block", spec.blocks, index, block
    if spec.slides:
        for slide in spec.slides:
            for index, block in enumerate(slide.blocks):
                yield "block", slide.blocks, index, block
    if spec.sheets:
        for index, sheet in enumerate(spec.sheets):
            yield "sheet", spec.sheets, index, sheet


def _slides_from_blocks(blocks: list[ContentBlock]) -> list[SpecSlide]:
    by_slide: dict[int, SpecSlide] = {}
    order: list[int] = []
    for block in blocks:
        slide_no = block.location.slide or 1
        if slide_no not in by_slide:
            by_slide[slide_no] = SpecSlide()
            order.append(slide_no)
        if block.type == BlockType.NOTE:
            by_slide[slide_no].notes = block.text
        else:
            by_slide[slide_no].blocks.append(content_to_spec_block(block))
    return [by_slide[number] for number in order]


def _strip_list_prefix(line: str) -> str:
    stripped = line.strip()
    return _LIST_PREFIX.sub("", stripped).strip() or stripped
